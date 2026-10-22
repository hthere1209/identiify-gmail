from pathlib import Path
from html import escape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, KeepTogether, HRFlowable
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
NAME = 'Jonathan Lee'
EMAIL = 'leejonathan793@gmail.com'
LINK = 'https://www.linkedin.com/in/jonathan-lee-926996443/'
NOTICE = 'REVIEW DRAFT — Proposed skills and responsibility bullets require your confirmation before submission.'
SUMMARY = ('Software and AI engineer with a bachelor’s degree in Computer Science from the National University of Singapore. '
           'Career progression from engineering intern to software developer, AI developer, and AI engineer, '
           'with roles at companies in Singapore, Poland, and the United States. '
           'Seeking AI training opportunities focused on coding, technical reasoning, and model-response evaluation.')
JOBS = [
 ('Allsage Tech', 'AI Engineer', 'United States', 'Feb 2025 – Present', [
  'Develop AI-enabled application components by translating project requirements into modular software, clearly defined interfaces, and testable workflows, with attention to maintainability, reliable integration, and expected system behavior.',
  'Evaluate model-generated responses against defined task requirements and quality criteria, assessing technical correctness, completeness, reasoning consistency, and instruction following while documenting specific errors and actionable feedback.',
  'Design structured test cases covering common user inputs, boundary conditions, and ambiguous requests; define expected outcomes and compare observed behavior to identify reliability gaps and support consistent evaluation across iterations.',
  'Investigate failures across input data, prompts, and application logic by reproducing problematic scenarios, tracing contributing factors, and documenting findings that help distinguish model limitations from software defects and guide targeted improvements.',
  'Collaborate with engineers to review evaluation findings, refine implementation decisions, and clarify acceptance criteria; maintain evaluation notes and technical documentation that support reproducible testing, knowledge sharing, and subsequent development.',
 ]),
 ('Veloce Software', 'AI Developer', 'Poland', 'Apr 2021 – Nov 2024', [
  'Develop AI-related application features and connect model outputs with downstream software workflows.',
  'Prepare and review datasets for missing values, inconsistent labels, duplicate records, and unsuitable examples.',
  'Compare model outputs with expected results to identify recurring errors and prioritize improvements.',
  'Test model integrations across representative scenarios, documenting limitations and reproducible failure cases.',
  'Work with software engineers to resolve integration issues and maintain implementation and testing documentation.',
 ]),
 ('Itpace', 'Software Developer', 'Singapore', 'Mar 2017 – Aug 2020', [
  'Develop and maintain application features from functional requirements, translating user needs into software behavior.',
  'Investigate defects by reproducing issues, tracing execution paths, and implementing targeted fixes.',
  'Write and maintain tests for core functionality, exceptional inputs, and regression prevention.',
  'Review code for correctness, readability, maintainability, and appropriate handling of edge cases.',
  'Collaborate with teammates to clarify requirements and document technical decisions and application behavior.',
 ]),
 ('Infini Technologies', 'Engineering Intern', 'Singapore', 'Nov 2015 – Dec 2016', [
  'Support engineers with development, testing, and maintenance tasks under technical supervision.',
  'Implement scoped features and bug fixes, incorporating feedback from code reviews.',
  'Execute test cases, compare observed behavior with requirements, and record reproducible defects.',
  'Prepare technical notes and progress updates to support project coordination and knowledge sharing.',
 ]),
]

# Use past tense for completed roles.
past = {'Develop': 'Developed', 'Prepare': 'Prepared', 'Compare': 'Compared', 'Test': 'Tested', 'Work': 'Worked', 'Investigate': 'Investigated', 'Write': 'Wrote', 'Review': 'Reviewed', 'Collaborate': 'Collaborated', 'Support': 'Supported', 'Implement': 'Implemented', 'Execute': 'Executed'}
for j in JOBS[1:]:
    for i, bullet in enumerate(j[4]):
        first, rest = bullet.split(' ', 1)
        j[4][i] = past.get(first, first) + ' ' + rest
# Match coordinated verbs to the completed-role tense.
for j in JOBS[1:]:
    j[4][:] = [b.replace('Developed and maintain ', 'Developed and maintained ').replace('Prepared and review ', 'Prepared and reviewed ').replace('Wrote and maintain ', 'Wrote and maintained ').replace('Executed test cases, compare ', 'Executed test cases, compared ').replace('requirements, and record ', 'requirements, and recorded ') for b in j[4]]

SKILLS = [
 ('Model evaluation', 'Response comparison, rubric-based assessment, instruction following, factual consistency.'),
 ('Coding and reasoning', 'Code review, debugging, test-case design, edge-case analysis, technical explanations.'),
 ('Training-data quality', 'Prompt and reference-answer writing, annotation consistency, data validation.'),
 ('Engineering practices', 'Requirements analysis, software testing, technical documentation, collaborative review.'),
]

pdfmetrics.registerFont(TTFont('Resume', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
pdfmetrics.registerFont(TTFont('ResumeBold', '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))
pdfmetrics.registerFontFamily('Resume', normal='Resume', bold='ResumeBold')
navy = colors.HexColor('#162D46'); blue=colors.HexColor('#315E7B'); ink=colors.HexColor('#344154')
styles={
 'name': ParagraphStyle('name',fontName='ResumeBold',fontSize=27,leading=32,textColor=navy,spaceAfter=4),
 'title': ParagraphStyle('title',fontName='Resume',fontSize=12,leading=17,textColor=blue,spaceAfter=7),
 'contact': ParagraphStyle('contact',fontName='Resume',fontSize=9,leading=14,textColor=ink,spaceAfter=7),
 'notice': ParagraphStyle('notice',fontName='Resume',fontSize=8,leading=12,textColor=colors.HexColor('#765B31'),backColor=colors.HexColor('#FAF5E9'),borderPadding=7,spaceBefore=6,spaceAfter=12),
 'section': ParagraphStyle('section',fontName='ResumeBold',fontSize=10,leading=15,textColor=blue,spaceBefore=12,spaceAfter=9),
 'body': ParagraphStyle('body',fontName='Resume',fontSize=9.2,leading=13.2,textColor=ink,spaceAfter=4),
 'role': ParagraphStyle('role',fontName='ResumeBold',fontSize=11,leading=16,textColor=navy,spaceAfter=3),
 'meta': ParagraphStyle('meta',fontName='Resume',fontSize=9,leading=14,textColor=blue,spaceAfter=6),
 'bullet': ParagraphStyle('bullet',fontName='Resume',fontSize=9.2,leading=13.2,textColor=ink,leftIndent=11,firstLineIndent=-10,spaceAfter=4),
}
def para(t,style='body'):return Paragraph(t,styles[style])
def job(j):
 company,role,loc,dates,bullets=j
 parts=[para(f'{escape(role)} | {escape(company)}','role'),para(f'{loc} · {dates}','meta')]
 parts += [para('• '+escape(b),'bullet') for b in bullets]
 parts.append(Spacer(1,6))
 return KeepTogether(parts)
def footer(canvas,doc):
 canvas.setStrokeColor(colors.HexColor('#DEE5EB'));canvas.line(43,40,A4[0]-43,40)
 canvas.setFont('Resume',8);canvas.setFillColor(colors.HexColor('#697789'))
 canvas.drawString(43,26,'Jonathan Lee • Review draft')
 canvas.drawRightString(A4[0]-43,26,f'{doc.page} / 2')
items=[para(NAME,'name'),para('SOFTWARE & AI ENGINEER','title'),
 para(f'Singapore · <link href="mailto:{EMAIL}" color="#315E7B">{EMAIL}</link><br/><link href="{LINK}" color="#315E7B">linkedin.com/in/jonathan-lee-926996443</link>','contact'),
 HRFlowable(width='100%',thickness=1.3,color=navy),para(NOTICE,'notice'),
 para('PROFESSIONAL PROFILE','section'),para(SUMMARY),
 para('PROFESSIONAL EXPERIENCE · PROPOSED RESPONSIBILITIES','section'),job(JOBS[0]),job(JOBS[1]),PageBreak(),
 para('Jonathan Lee','role'),para('Software & AI Engineer · Experience continued','contact'),HRFlowable(width='100%',thickness=1.3,color=navy),
 para(NOTICE,'notice'),job(JOBS[2]),job(JOBS[3]),
 para('AI TRAINING SKILLS · PROPOSED FOR REVIEW','section')]
for title,body in SKILLS:items.append(para(f'<b>{escape(title)}:</b> {escape(body)}'))
items += [para('EDUCATION','section'),para('<b>National University of Singapore</b>'),para('Bachelor’s degree in Computer Science · Graduated 2015')]
SimpleDocTemplate(str(ROOT/'Jonathan_Lee_Resume_Draft.pdf'),pagesize=A4,leftMargin=43,rightMargin=43,topMargin=38,bottomMargin=53,title='Jonathan Lee — Software & AI Engineer (Review Draft)',author=NAME).build(items,onFirstPage=footer,onLaterPages=footer)

# Editable Word version, using the same content and reading order.
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.52);sec.bottom_margin=Inches(.60);sec.left_margin=sec.right_margin=Inches(.60)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(9.5)
normal.font.color.rgb=RGBColor.from_string('344154');normal.paragraph_format.space_after=Pt(4)
normal.paragraph_format.line_spacing=1.0
for style in ['Heading 1','Heading 2']:
 doc.styles[style].font.name='Calibri';doc.styles[style].font.color.rgb=RGBColor.from_string('315E7B')
 doc.styles[style].font.size=Pt(11)
 doc.styles[style].paragraph_format.space_before=Pt(10);doc.styles[style].paragraph_format.space_after=Pt(6)
def dp(text,bold=False,size=None):
 p=doc.add_paragraph();r=p.add_run(text);r.bold=bold
 if size:r.font.size=Pt(size)
 return p
def link(p,text,url):
 rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel);r=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'315E7B');props.append(color);r.append(props);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)
def djob(j):
 company,role,loc,dates,bullets=j
 p=dp(f'{role} | {company}',True,11);p.paragraph_format.keep_with_next=True
 p=dp(f'{loc} · {dates}');p.paragraph_format.keep_with_next=True
 for b in bullets:
  p=doc.add_paragraph(b,style='List Bullet');p.paragraph_format.space_after=Pt(4)
  p.paragraph_format.keep_together=True
p=dp(NAME,True,27);dp('SOFTWARE & AI ENGINEER',True,12)
p=dp('Singapore · ');link(p,EMAIL,'mailto:'+EMAIL)
p=doc.add_paragraph();link(p,'linkedin.com/in/jonathan-lee-926996443',LINK)
dp(NOTICE,False,8)
doc.add_heading('PROFESSIONAL PROFILE',1);dp(SUMMARY)
doc.add_heading('PROFESSIONAL EXPERIENCE · PROPOSED RESPONSIBILITIES',1)
djob(JOBS[0]);djob(JOBS[1]);doc.add_page_break()
dp('Jonathan Lee | Experience continued',True,11);dp(NOTICE,False,8)
djob(JOBS[2]);djob(JOBS[3]);doc.add_heading('AI TRAINING SKILLS · PROPOSED FOR REVIEW',1)
for title,body in SKILLS:
 p=doc.add_paragraph();p.add_run(title+': ').bold=True;p.add_run(body)
doc.add_heading('EDUCATION',1);dp('National University of Singapore',True);dp('Bachelor’s degree in Computer Science · Graduated 2015')
f=sec.footer.paragraphs[0];f.add_run('Jonathan Lee • Review draft  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
doc.core_properties.title='Jonathan Lee — Software & AI Engineer (Review Draft)';doc.core_properties.author=NAME
doc.save(ROOT/'Jonathan_Lee_Resume_Draft.docx')
print('Created PDF and editable DOCX review drafts.')
