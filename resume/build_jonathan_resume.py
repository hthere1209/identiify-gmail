from pathlib import Path
from html import escape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
EMAIL='leejonathan793@gmail.com'
LINK='https://www.linkedin.com/in/jonathan-lee-926996443/'
SUMMARY=[
 'Software and AI Engineer with more than 8 years of experience building, deploying, and maintaining applications, backend services, APIs, and data-driven systems.',
 'Strong background in AI model evaluation, mathematical reasoning, and LLM training data, including benchmarking, RLHF/SFT annotation, and model assessment. Uses Python, machine learning tools, and data analysis to evaluate responses, identify reasoning flaws, and develop practical improvements.',
 'Translates business requirements into working software and collaborates with engineering teams throughout development, testing, and delivery.'
]
JOBS=[
 ('AI Engineer','AllsageTech','United States','Feb 2025 – Present',[
  'Build AI features that help users complete tasks more effectively, working with the team to understand user needs and develop software that meets their expectations.',
  'Review AI-generated responses for factual accuracy, correctness, completeness, clarity, and alignment with user requirements; analyze errors and recommend improvements.',
  'Develop questions and realistic test scenarios, including ambiguous requests and invalid inputs, to assess AI performance and identify areas requiring improvement.',
  'Investigate issues with AI features, collaborate with developers to identify and resolve their causes, and verify that changes restore expected behavior.',
  'Work with developers and stakeholders to establish quality standards, clarify expectations, communicate progress, and document features and known limitations.'
 ]),
 ('AI Developer','Veloce Software','Poland','Apr 2021 – Nov 2024',[
  'Developed AI features and integrated them into existing applications so users could access AI capabilities within their everyday workflows.',
  'Prepared and reviewed AI training data, identifying missing information, irrelevant examples, duplicate records, and errors to improve data quality.',
  'Compared AI-generated results with expected outcomes to identify recurring weaknesses and help prioritize improvements.',
  'Tested AI features in realistic application scenarios, assessed their strengths and limitations, and provided developers with clear examples of problems to investigate.',
  'Collaborated with software developers to resolve issues between AI features and other application components, documenting workflows, implementation details, and testing procedures.'
 ]),
 ('Software Developer','Itpace','Singapore','Mar 2017 – Aug 2020',[
  'Developed and maintained application features based on user needs, working with the team to understand product usage and translate requirements into practical functions supporting daily activities.',
  'Investigated software issues by reproducing bugs, identifying root causes, and implementing fixes. Verified that affected features worked correctly without disrupting other parts of the application.',
  'Created and updated tests for critical features under normal and unexpected conditions, helping detect problems during development and confirm that existing functionality remained reliable after changes.',
  'Reviewed teammates’ code, identified improvements, and applied review feedback in line with team standards to simplify future maintenance and reduce the risk of defects.',
  'Collaborated with teammates to clarify requirements, discuss technical options, and share progress. Maintained documentation and development records to support everyday use and future improvements.'
 ]),
 ('Engineering Intern','Infini Technologies','Singapore','Nov 2015 – Dec 2016',[
  'Supported experienced engineers with software development, testing, and maintenance, learning how requirements became working features and completing assigned tasks under senior guidance.',
  'Built small features and fixed straightforward software issues, following senior developers’ instructions and applying code-review feedback to improve quality, documentation, and consistency.',
  'Tested application features by following defined steps and comparing actual results with expected outcomes. Documented issues clearly so developers could reproduce, investigate, and resolve them.',
  'Prepared technical notes and progress reports covering completed tasks, questions, and outstanding issues. Organized this information to support peer review and subsequent development work.'
 ])]
SKILLS=[
 ('AI Training & Model Evaluation','LLM response evaluation, response ranking, rubric-based assessment, instruction following, factual accuracy review, reasoning assessment, coding evaluation, error analysis.'),
 ('Training Data & Quality Assurance','Prompt writing, reference-answer development, data annotation and review, SFT data curation, preference data for RLHF, data validation, deduplication, quality guidelines.'),
 ('Programming & Problem Solving','Python, JavaScript, TypeScript, SQL, data structures, algorithms, object-oriented programming, debugging, code review, test-case design.'),
 ('AI, Machine Learning & Data','Generative AI, LLM integration, prompt engineering, AI API integration, machine learning fundamentals, data preprocessing and analysis, NumPy, Pandas, SciPy, TensorFlow.'),
 ('Backend & Application Development','Node.js, Express.js, Django, FastAPI, REST APIs, API integration, webhooks, database integration, application architecture, reusable software components.'),
 ('Frontend Development','React, Next.js, HTML5, CSS3, responsive web design, Tailwind CSS, Bootstrap.'),
 ('Databases & Cloud','PostgreSQL, MySQL, MongoDB, Redis, Firebase, SQL/NoSQL data modeling, query optimization, AWS, Docker, CI/CD.'),
 ('Testing & Engineering Practices','Unit, integration, API, and regression testing; Git, GitHub, Agile development, requirements analysis, technical documentation, collaborative code review.')]

# A single-column layout keeps the reading order predictable in PDF extraction.
pdfmetrics.registerFont(TTFont('Resume','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
pdfmetrics.registerFont(TTFont('ResumeBold','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))
pdfmetrics.registerFontFamily('Resume',normal='Resume',bold='ResumeBold')
NAVY=colors.HexColor('#17334B');BLUE=colors.HexColor('#32677A');INK=colors.HexColor('#344154')
styles={
 'name':ParagraphStyle('name',fontName='ResumeBold',fontSize=28,leading=33,textColor=NAVY,spaceAfter=4),
 'title':ParagraphStyle('title',fontName='Resume',fontSize=11,leading=16,textColor=BLUE,spaceAfter=6),
 'contact':ParagraphStyle('contact',fontName='Resume',fontSize=8.6,leading=13,textColor=INK,spaceAfter=10),
 'section':ParagraphStyle('section',fontName='ResumeBold',fontSize=12,leading=17,textColor=BLUE,spaceBefore=9,spaceAfter=6,keepWithNext=True),
 'body':ParagraphStyle('body',fontName='Resume',fontSize=9.1,leading=13,textColor=INK,spaceAfter=5),
 'role':ParagraphStyle('role',fontName='ResumeBold',fontSize=10.5,leading=15,textColor=NAVY,spaceAfter=2),
 'meta':ParagraphStyle('meta',fontName='Resume',fontSize=8.7,leading=12.5,textColor=BLUE,spaceAfter=5),
 'bullet':ParagraphStyle('bullet',fontName='Resume',fontSize=9.1,leading=12.4,textColor=INK,leftIndent=10,firstLineIndent=-10,spaceAfter=3),
 'skill':ParagraphStyle('skill',fontName='Resume',fontSize=8.8,leading=11.8,textColor=INK,spaceAfter=4),
}
def p(t,style='body'):return Paragraph(t,styles[style])
def section(t):return p(t,'section')
def job(j):
 role,company,loc,dates,bullets=j
 heading=p(escape(role)+' | '+escape(company),'role'); heading.keepWithNext=True
 metadata=p(loc+' · '+dates,'meta'); metadata.keepWithNext=True
 return [heading,metadata]+[p('• '+escape(b),'bullet') for b in bullets]+[Spacer(1,6)]
def footer(c,d):
 c.setStrokeColor(colors.HexColor('#D6E1E7'));c.line(42,37,A4[0]-42,37)
 c.setFont('Resume',7.5);c.setFillColor(colors.HexColor('#6E7C88'))
 c.drawString(42,24,'Jonathan Lee · Software & AI Engineer');c.drawRightString(A4[0]-42,24,str(d.page))
items=[p('Jonathan Lee','name'),p('SOFTWARE & AI ENGINEER','title'),
 p(f'Singapore · <link href="mailto:{EMAIL}" color="#32677A">{EMAIL}</link><br/><link href="{LINK}" color="#32677A">linkedin.com/in/jonathan-lee-926996443</link>','contact'),
 HRFlowable(width='100%',thickness=1.4,color=NAVY),section('PROFESSIONAL SUMMARY')]
items += [p(escape(t)) for t in SUMMARY]
items.append(section('PROFESSIONAL EXPERIENCE'))
for entry in JOBS: items.extend(job(entry))
items += [section('EDUCATION'),p('<b>National University of Singapore</b>'),p('Bachelor’s degree in Computer Science · January 2015'),section('TECHNICAL SKILLS')]
for title,body in SKILLS:items.append(p('<b>'+escape(title)+':</b> '+escape(body),'skill'))
SimpleDocTemplate(str(ROOT/'Jonathan_Lee_Resume.pdf'),pagesize=A4,leftMargin=42,rightMargin=42,topMargin=30,bottomMargin=46,title='Jonathan Lee | Software & AI Engineer',author='Jonathan Lee').build(items,onFirstPage=footer,onLaterPages=footer)

# Word source: matching content, editable paragraphs, active hyperlinks.
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.48);sec.bottom_margin=Inches(.55);sec.left_margin=sec.right_margin=Inches(.58)
style=doc.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(9.5);style.font.color.rgb=RGBColor.from_string('344154')
style.paragraph_format.line_spacing=1;style.paragraph_format.space_after=Pt(4)
for name in ['Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(12);st.font.color.rgb=RGBColor.from_string('32677A')
 st.paragraph_format.space_before=Pt(9);st.paragraph_format.space_after=Pt(5)
def dp(t,bold=False,size=None):
 para=doc.add_paragraph();r=para.add_run(t);r.bold=bold
 if size:r.font.size=Pt(size)
 return para
def hyperlink(para,text,url):
 rel=para.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel);run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'32677A');props.append(color);run.append(props);t=OxmlElement('w:t');t.text=text;run.append(t);h.append(run);para._p.append(h)
def djob(j):
 role,company,loc,dates,bullets=j
 para=dp(role+' | '+company,True,11);para.paragraph_format.keep_with_next=True
 para=dp(loc+' · '+dates);para.paragraph_format.keep_with_next=True
 for b in bullets:
  para=doc.add_paragraph(b,style='List Bullet');para.paragraph_format.keep_together=True;para.paragraph_format.space_after=Pt(4)
dp('Jonathan Lee',True,28);dp('SOFTWARE & AI ENGINEER',True,11)
para=dp('Singapore · ');hyperlink(para,EMAIL,'mailto:'+EMAIL)
para=doc.add_paragraph();hyperlink(para,'linkedin.com/in/jonathan-lee-926996443',LINK)
doc.add_heading('PROFESSIONAL SUMMARY',1)
for t in SUMMARY:dp(t)
doc.add_heading('PROFESSIONAL EXPERIENCE',1)
djob(JOBS[0]);djob(JOBS[1])
djob(JOBS[2]);djob(JOBS[3]);doc.add_heading('EDUCATION',1)
dp('National University of Singapore',True);dp('Bachelor’s degree in Computer Science · January 2015')
doc.add_heading('TECHNICAL SKILLS',1)
for title,body in SKILLS:
 para=doc.add_paragraph();para.add_run(title+': ').bold=True;para.add_run(body)
f=sec.footer.paragraphs[0];f.add_run('Jonathan Lee · Software & AI Engineer  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
doc.core_properties.title='Jonathan Lee | Software & AI Engineer';doc.core_properties.author='Jonathan Lee'
doc.save(ROOT/'Jonathan_Lee_Resume.docx')
print('Created Jonathan_Lee_Resume.pdf and Jonathan_Lee_Resume.docx')
